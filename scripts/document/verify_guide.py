"""Check the consultant pack against the operative rubric and rendered PDF."""
import json
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'output/documents/meeting-coach-scoring-guide'
rubric = json.loads((ROOT / 'lib/rubric.json').read_text())


def normal(text):
    # Preserve real hyphens while joining PDF lines wrapped after a hyphen.
    return re.sub(r'\s+', ' ', re.sub(r'-\s*\n\s*', '-', text)).strip()


doc = Document(str(BASE.with_suffix('.docx')))
doc_text = '\n'.join(p.text for p in doc.paragraphs)
doc_text += '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
pdf_text = subprocess.check_output(['pdftotext', '-layout', str(BASE.with_suffix('.pdf')), '-'], text=True)
pages = [p for p in pdf_text.split('\f') if p.strip()]
assert len(pages) == 9, f'Expected 9 pages, got {len(pages)}'

for index, section in enumerate(rubric['scores']):
    page = normal(pages[index + 2])
    assert section['name'].replace('&', 'and') in page
    for wording in [c['text'] for c in section['criteria']] + list(section['anchors'].values()):
        assert normal(wording) in page, f'Missing on section page: {wording}'
        assert wording in doc_text, f'Missing from DOCX: {wording}'
    assert 'Point for consultant review' in page
assert 'For Afrikaans-first clients' not in pdf_text + doc_text
assert '21 ÷ 6 = 3.50' in normal(pdf_text)

with tempfile.TemporaryDirectory(prefix='meeting-coach-document-qa-') as directory:
    bbox = Path(directory) / 'bbox.html'
    subprocess.run(['pdftotext', '-bbox', str(BASE.with_suffix('.pdf')), str(bbox)], check=True)
    namespace = {'x': 'http://www.w3.org/1999/xhtml'}
    for index, page in enumerate(ET.parse(bbox).findall('.//x:page', namespace)):
        words = page.findall('x:word', namespace)
        footer_index = next(i for i, w in enumerate(words) if w.text == 'Rubric')
        footer = words[footer_index]
        footer_y = float(footer.get('yMin'))
        body = [w for w in words[:footer_index] if float(w.get('yMin')) > 40]
        assert body
        assert max(float(w.get('yMax')) for w in body) < footer_y - 20, f'Footer collision on page {index + 1}'

print('PASS: 9 pages; all 7 sections, 24 criteria and 21 anchors on their intended pages; exact DOCX content; overall calculation; footer separation.')
print('Visual QA: inspect rendered pages separately; this check does not assess design or scoring validity.')
