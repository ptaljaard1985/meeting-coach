import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
rubric=json.loads((ROOT/'lib/rubric.json').read_text())
notes=json.loads((ROOT/'scripts/document/section-notes.json').read_text())
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.28)
for name in ['Normal','Body Text']:
 st=doc.styles[name];st.font.name='Aptos';st.font.size=Pt(10.5);st.font.color.rgb=RGBColor.from_string('222A2A');st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.06
for name,size,font in [('Title',29,'Georgia'),('Subtitle',12,'Aptos'),('Heading 1',22,'Georgia'),('Heading 2',12,'Aptos'),('Heading 3',10.5,'Aptos')]:
 st=doc.styles[name];st.font.name=font;st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=name in ['Heading 2','Heading 3'];st.paragraph_format.space_before=Pt(7 if name!='Title' else 0);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
 for attr in ['ascii','hAnsi','cs','eastAsia']:st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+attr),font)
 for attr in ['asciiTheme','hAnsiTheme','csTheme','eastAsiaTheme']:st.element.get_or_add_rPr().get_or_add_rFonts().attrib.pop(qn('w:'+attr),None)
for st in doc.styles:
 if st.type==1:
  el=st.element
  for e in list(el.iter(qn('w:pBdr'))):e.getparent().remove(e)
header=sec.header.paragraphs[0];header.text='MEETING COACH  /  SCORING GUIDE';header.style='Normal'
for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Rubric 1.1  |  2 October 2026  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def para(text='',style=None,boldlead=None):
 p=doc.add_paragraph(style=style)
 if boldlead and text.startswith(boldlead):p.add_run(boldlead).bold=True;p.add_run(text[len(boldlead):])
 else:p.add_run(text)
 return p

def heading(text,level=2):return doc.add_heading(text,level)

def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,text in enumerate(headers):t.rows[0].cells[i].text=text
 for row in rows:
  cells=t.add_row().cells
  for i,text in enumerate(row):cells[i].text=str(text)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders);mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   tcPr.append(mar);fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'344D5C' if ri==0 else ('F1F4F5' if ri%2 else 'FFFFFF'));tcPr.append(fill)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.04
    for r in p.runs:r.font.size=Pt(9.5);r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '222A2A');r.bold=ri==0
 return t

para('Meeting Coach Scoring Guide','Title')
para('Seven sections · 24 criteria · Adviser communication','Subtitle')
para('Prepared for consultant review on 2 October 2026')
para('This guide explains what Meeting Coach assesses, how evidence supports a judgment, and how the resulting scores are calculated. Please review whether the criteria and anchors are fair, observable and useful across different meeting types.')
para('The tool assesses the adviser\'s communication and meeting craft. It does not judge clients, evaluate investment suitability, certify compliance or generate client advice. English, Afrikaans and mixed-language meetings use the same criteria; coaching is in British English and quotes retain the original speech.')
heading('The seven sections')
table(['Section','Focus','Page'],[(str(i+1)+' '+s['name'],n['purpose'].split('. ')[0].rstrip('. ')+'.',str(i+3)) for i,(s,n) in enumerate(zip(rubric['scores'],notes))],[2.02,4.22,.55])
heading('How to read each section')
para('Each section has its own page with the current anchors and exact criterion wording. Explanations, evidence guidance and not-assessable examples are interpretation notes for consultant review; they do not add new scoring rules.')
heading('What a criterion verdict means')
table(['Verdict','Meaning'],[
 ('Met','The relevant behaviour is demonstrated by the cited interaction.'),
 ('Partially met','The behaviour is incomplete, inconsistent or only partly demonstrated.'),
 ('Not met','A real opportunity arose and was missed or contradicted; the relevant interaction must be cited.'),
 ('Not assessable','No relevant opportunity or adequate evidence is visible. This is not a failure.')],[1.38,5.41])
para('The Afrikaans-specific Plain language criterion has been removed. All seven sections remain; Plain language now has two criteria. This is the operative rubric, version 1.1.')

doc.add_page_break();heading('How scoring is determined',1)
para('The computer calculates conversation metrics. The AI makes evidence-based judgments against the rubric. The server validates the output and calculates the overall; it does not determine whether the AI\'s interpretation is fair.')
heading('From transcript to category score')
steps=[
 ('1 Confirm context','Select the adviser and confirm other speakers as clients or observers. Both voices applies only to a confirmed couple with two clients. Unsupported attribution is rejected.'),
 ('2 Calculate the metrics','Code calculates the figures below. Observer speech remains in context but is excluded from adviser/client metrics; it still breaks uninterrupted spans.'),
 ('3 Assess every criterion','Each verdict has a rationale. An assessable criterion requires a verbatim quote from a transcript segment; that segment supplies its speaker and any real timestamp.'),
 ('4 Apply the category anchors','The AI assigns each assessable category a value from 1 to 5 against the low, middle and high anchors. Middle means ordinary competence. Repeated strong evidence supports the top; clear counter-evidence supports the bottom.'),
 ('5 Validate and average','Section and criterion IDs must be complete and unique. Unassessable sections have no numeric value. The server calculates the equal-weight mean of assessable sections.')]
for title,text in steps:heading(title,3);para(text)
para('There is no fixed conversion from criterion verdicts to points, no arithmetic average of the criteria, and no special weighting for sections with more criteria. Values 2 and 4, or decimals, are model judgments between the supplied anchors; all five levels are not separately defined in the current rubric.')
heading('What the calculated metrics mean')
table(['Metric','Calculation and limitation'],[
 ('Talk share','Estimated adviser duration ÷ (adviser + client duration). Use timestamp gaps when all starts are valid and non-decreasing; otherwise words at 140 words per minute.'),
 ('Question turns','Adviser speech segments ending in a question mark. Not a comprehensive linguistic inventory.'),
 ('Uninterrupted spans','Consecutive segments by the same speaker. Observer speech breaks a run. Gaps include pauses; repeated starts and an unknown final tail use word estimates.'),
 ('Duration','Known total if supplied; otherwise the visible timestamp span plus estimated tail, or a word-based estimate. Known duration never creates quote timestamps.'),
 ('Client balance','For a couple only: larger client speaking duration ÷ smaller. Approximate 2:1 and 3:1 thresholds guide judgment; they do not measure understanding.')],[1.24,5.55])
heading('The overall score')
para('Example: six assessable category scores of 4, 3, 3, 4, 4 and 3 sum to 21. Both voices is n/a in a solo meeting. Overall = 21 ÷ 6 = 3.50 out of 5. A fully unassessable review has overall n/a. Displayed precision is not proof of scientific accuracy.')
para('A valid quote does not prove a fair interpretation. Scores can vary between AI runs. Calibration should check evidence coverage, fair n/a decisions, useful coaching and agreement with human judgments.')

for i,(score,note) in enumerate(zip(rubric['scores'],notes)):
 doc.add_page_break();heading(str(i+1)+' '+score['name'].replace('&','and'),1)
 para(note['purpose'])
 heading('Current scoring anchors')
 table(['Anchor','Current rubric wording'],[(label,score['anchors'][key]) for label,key in [('Low','low'),('Middle','mid'),('High','high')]],[.85,5.94])
 para(note['scoring'])
 for j,(criterion,detail) in enumerate(zip(score['criteria'],note['criteria'])):
  heading(f'{i+1}.{j+1} '+detail['heading'],3)
  p=para(criterion['text']);p.paragraph_format.space_after=Pt(3)
  for r in p.runs:r.italic=True;r.font.size=Pt(10)
  para(detail['meaning'])
  para('Evidence to look for: '+detail['look'],boldlead='Evidence to look for: ')
  para('Not assessable when: '+detail['na'],boldlead='Not assessable when: ')
 heading('Point for consultant review',3);para(note['review'])

props=doc.core_properties;props.title='Meeting Coach Scoring Guide';props.subject='Seven section rubric and scoring method';props.author='Meeting Coach';props.keywords='meeting craft, consultant review, scoring';props.comments=''
output=ROOT/'output/documents/meeting-coach-scoring-guide.docx';doc.save(output)
print(output)
print('Sections:',len(rubric['scores']),'Criteria:',sum(len(s['criteria']) for s in rubric['scores']))
